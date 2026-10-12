"""섹터 ETF 보유 종목 — 「국가별 섹터종목」(사용자 2026-10-12)의 종목 목록.

Breadth 전략이 쓰는 섹터 ETF(`market_timing._BREADTH_SECTORS` — 미국 SPDR 11 · 한국 KODEX 13)
마다 **보유 비중 상위 15종목**을 매일 한 번 받는다.

- 미국 SPDR — 운용사(State Street) 일일 보유 엑셀 `holdings-daily-us-en-{티커}.xlsx`
  (주소 형식은 tidyquant `tq_stock_list.R` 가 2024-08-16 부터 쓰는 그것, 2026-10-12 확인).
  현금·선물 행(섹터 'Unassigned'·티커 없음)은 뺀다.
- 한국 KODEX — KRX ETF 구성종목(PDF, pykrx `get_etf_portfolio_deposit_file`). KRX 로그인
  (`KRX_ID`·`KRX_PW`)이 있어야 한다. 6자리 종목코드가 아닌 행(현금·선물)은 뺀다.

이 환경(샌드박스)에선 두 원천 다 닿지 못해 **실측하지 못했다** — 실패는 ETF 마다 사유를 남기고
화면이 그 ETF 칸에 적는다(#43·#82). 받은 목록은 하루 캐시하고, 오늘 실패하면 마지막으로 받은
목록을 **그 기준일과 함께** 보여 준다(낡았다는 사실은 화면이 말한다, #41).
"""
from __future__ import annotations

import io
import json
import logging
import math
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

log = logging.getLogger("bot.etf_holdings")

_KST = timezone(timedelta(hours=9))
_DIR = Path(os.environ.get("HOME") or str(Path.home())) / ".tradingagents" / "world_sectors" / "holdings"
TOP_N = 15
SSGA_URL = ("https://www.ssga.com/us/en/institutional/library-content/products/fund-data/"
            "etfs/us/holdings-daily-us-en-{t}.xlsx")
_TIMEOUT = (10, 30)
_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
       "(KHTML, like Gecko) Chrome/120.0 Safari/537.36")
_KR_CODE = re.compile(r"^\d{6}$")


def _finite(v) -> bool:
    """비중이 쓸 수 있는 수인가 — NaN·무한은 `json.dumps` 가 `NaN` 으로 써 브라우저 `JSON.parse` 가
    실패하고, 그 뒤 ETF 표가 전부 안 그려진다. 비중이 없는(None) 행은 남긴다(이름은 보인다)."""
    if v is None:
        return True
    try:
        return math.isfinite(float(v))
    except (TypeError, ValueError):
        return False


class HoldingsError(Exception):
    def __init__(self, reason: str):
        super().__init__(reason)
        self.reason = reason


def etfs(market: str) -> dict[str, str]:
    """Breadth 전략의 섹터 ETF — {티커: 섹터 이름}. 단일 출처(`market_timing`)를 그대로 읽는다(#38)."""
    from bot.market_timing import _BREADTH_SECTORS
    return dict(_BREADTH_SECTORS.get(market) or {})


def etf_label(market: str) -> str:
    from bot.market_timing import _BREADTH_SOURCE_LABEL
    return _BREADTH_SOURCE_LABEL.get(market, "")


# ── 미국 SPDR(State Street 엑셀) ─────────────────────────────────────────────
def parse_ssga_xlsx(data: bytes, top: int = TOP_N) -> tuple[list[dict], str]:
    """엑셀 바이트 → (상위 보유 [{code, name, weight}], 기준일 문자열). 순수.

    머리 몇 줄(펀드 이름·'Holdings: As of …')이 표 위에 있다 — 'Ticker' 와 'Weight' 가 같이 있는
    행을 머리행으로 찾는다(위치로 자르지 않는다, #46). 현금·선물 행은 섹터 'Unassigned' 이거나
    티커가 비어 있다."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    asof = ""
    head = None
    out: list[dict] = []
    for row in ws.iter_rows(values_only=True):
        cells = [("" if c is None else c) for c in row]
        texts = [str(c).strip() for c in cells]
        if head is None:
            for t in texts:
                m = re.search(r"as of[:\s]+(.+)$", t, re.I)
                if m:
                    asof = m.group(1).strip()
            low = [t.lower() for t in texts]
            if "ticker" in low and "weight" in low:
                head = {t: i for i, t in enumerate(low) if t}
            continue
        if not any(texts):
            if out:
                break
            continue
        tk = texts[head["ticker"]] if head.get("ticker") is not None else ""
        name = texts[head["name"]] if head.get("name") is not None else tk
        sector = texts[head["sector"]] if head.get("sector") is not None else ""
        try:
            w = float(cells[head["weight"]])
        except (TypeError, ValueError):
            continue
        if not math.isfinite(w):
            continue
        if not tk or tk in ("-", "--") or sector.lower() in ("unassigned", "-"):
            continue
        out.append({"code": tk.upper(), "name": name, "weight": w})
    if head is None:
        raise HoldingsError("엑셀에서 'Ticker'·'Weight' 머리행을 못 찾음 — 파일 형식이 바뀐 것으로 보임")
    out.sort(key=lambda r: -r["weight"])
    return out[:top], asof


def fetch_us(ticker: str, session=None) -> tuple[list[dict], str]:
    import requests
    from bot.naver_diag import http_reason
    s = session or requests
    url = SSGA_URL.format(t=ticker.lower())
    try:
        r = s.get(url, headers={"user-agent": _UA}, timeout=_TIMEOUT)
    except Exception as exc:                                   # noqa: BLE001
        raise HoldingsError(http_reason(None, exc=exc)) from exc
    if getattr(r, "status_code", None) != 200:
        raise HoldingsError(http_reason(getattr(r, "status_code", None),
                                        body=getattr(r, "content", b"")))
    content = getattr(r, "content", b"") or b""
    if not content.startswith(b"PK"):
        raise HoldingsError("엑셀이 아닌 본문이 옴(로그인·차단 페이지로 보임)")
    rows, asof = parse_ssga_xlsx(content)
    if not rows:
        raise HoldingsError("엑셀에 보유 종목 행이 없음")
    return rows, asof


# ── 한국 KODEX(KRX PDF) ──────────────────────────────────────────────────────
def parse_kr_pdf(df, top: int = TOP_N) -> list[dict]:
    """pykrx PDF DataFrame → 상위 보유(순수). 비중 열이 있으면 비중, 없으면 평가금액으로 줄 세운다.
    6자리 종목코드가 아닌 행(현금·선물·채권)은 뺀다."""
    if df is None or getattr(df, "empty", True):
        return []
    cols = list(df.columns)

    def col(*names):
        for n in names:
            if n in cols:
                return n
        return None

    wcol, vcol = col("비중", "구성비중"), col("금액", "평가금액", "시가총액")
    ncol = col("구성종목명", "종목명")
    out = []
    for code, r in df.iterrows():
        c = str(code).strip()
        if not _KR_CODE.match(c):
            continue
        w = float(r[wcol]) if wcol and r[wcol] == r[wcol] else None
        v = float(r[vcol]) if vcol and r[vcol] == r[vcol] else None
        w = w if w is not None and math.isfinite(w) else None
        v = v if v is not None and math.isfinite(v) else None
        out.append({"code": c, "name": str(r[ncol]) if ncol else c, "weight": w, "_v": v})
    tot = sum(x["_v"] for x in out if x["_v"])
    for x in out:
        if (x["weight"] is None or x["weight"] == 0) and x["_v"] and tot > 0:
            x["weight"] = x["_v"] / tot * 100
    out.sort(key=lambda x: -(x["weight"] or 0))
    for x in out:
        x.pop("_v", None)
    return out[:top]


def fetch_kr(ticker: str, date: str) -> tuple[list[dict], str]:
    """KRX PDF — `date` = 그 거래일(YYYY-MM-DD). KRX 로그인이 없으면 실패(사유를 남긴다)."""
    try:
        from bot.pykrx_client import krx_login_ready
        if not krx_login_ready():
            from bot.env_keys import env_diag
            # 갈래(파일 못 찾음 · 키 없음 · 값 비었음 …)를 같이 적는다 — '미설정' 만 적으면
            # 이미 있는 키를 넣으러 간다(#82·#294). 값은 안 찍는다.
            raise HoldingsError("KRX_ID/KRX_PW 미설정 — KRX ETF 구성종목을 받을 수 없음 "
                                f"({env_diag('KRX_ID', 'KRX_PW') or '갈래 판정 불가'})")
        from pykrx import stock
    except HoldingsError:
        raise
    except Exception as exc:                                   # noqa: BLE001
        raise HoldingsError(f"pykrx 없음({type(exc).__name__})") from exc
    code = ticker.split(".")[0]
    try:
        df = stock.get_etf_portfolio_deposit_file(code, date.replace("-", ""))
    except Exception as exc:                                   # noqa: BLE001
        raise HoldingsError(f"KRX ETF 구성종목 조회 실패({type(exc).__name__}: {str(exc)[:80]})") from exc
    rows = parse_kr_pdf(df)
    if not rows:
        raise HoldingsError(f"KRX 가 {date} 구성종목을 주지 않음(빈 표)")
    return rows, date


# ── 캐시 ───────────────────────────────────────────────────────────────────
def _path(ticker: str) -> Path:
    return _DIR / f"{ticker.replace('/', '_')}.json"


def load(ticker: str) -> dict | None:
    try:
        d = json.loads(_path(ticker).read_text(encoding="utf-8"))
    except Exception:                                          # noqa: BLE001
        return None
    return d if isinstance(d, dict) else None


def _save(ticker: str, obj: dict) -> None:
    p = _path(ticker)
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_name(f"{p.name}.tmp{os.getpid()}")
    tmp.write_text(json.dumps(obj, ensure_ascii=False), encoding="utf-8")
    os.replace(tmp, p)


RETRY_HOURS = 2
MAX_TRIES = 4


def _retry_due(cur: dict, today: str, now: datetime) -> bool:
    """오늘 실패한 ETF 를 다시 물을 때인가 — 일시 장애 한 번이 하루를 비우지 않게(#161) 하되,
    원천이 오래 죽었을 때 15분마다 두드리지도 않게(#116) `RETRY_HOURS` 간격 · 하루 `MAX_TRIES` 회."""
    if cur.get("tried_day") != today:
        return True
    if int(cur.get("tries") or 1) >= MAX_TRIES:
        return False
    last = _tried_at(cur)
    return last is None or now - last >= timedelta(hours=RETRY_HOURS)


def _tried_at(cur: dict) -> datetime | None:
    try:
        last = datetime.fromisoformat(str(cur.get("tried_at")))
    except (TypeError, ValueError):
        return None
    return last if last.tzinfo else last.replace(tzinfo=_KST)


def retry_note(h: dict, now: datetime) -> str:
    """실패한 보유 목록을 언제 다시 묻나 — 화면이 '기다리면 되나' 를 사실대로 말하게(#380)."""
    today = now.astimezone(_KST).date().isoformat()
    if h.get("tried_day") != today:
        return "다음 주기에 다시 묻습니다"
    n = int(h.get("tries") or 1)
    if n >= MAX_TRIES:
        return f"오늘 {n}회 실패 — 내일 다시 묻습니다"
    last = _tried_at(h)
    nxt = f" · 다음 시도 {(last + timedelta(hours=RETRY_HOURS)).astimezone(_KST):%H:%M}" if last else ""
    return f"오늘 {n}/{MAX_TRIES}회 실패{nxt}"


def refresh(market: str, now: datetime, kr_date: str | None = None) -> dict:
    """하루 한 번 — 오늘(KST) 받은 목록이 있으면 건너뛴다. 실패하면 옛 목록은 두고
    `error`·`tried` 만 갈아 적는다(화면이 '마지막으로 받은 목록 · 오늘 실패 사유' 를 말한다).
    실패한 날은 `RETRY_HOURS` 마다 다시 묻는다(하루 `MAX_TRIES` 회까지)."""
    today = now.astimezone(_KST).date().isoformat()
    res = {"ok": [], "failed": []}
    for t in etfs(market):
        cur = load(t) or {}
        if cur.get("fetched_day") == today or not _retry_due(cur, today, now):
            continue
        tries = int(cur.get("tries") or 1) + 1 if cur.get("tried_day") == today else 1
        try:
            if market == "US":
                rows, asof = fetch_us(t)
            else:
                if not kr_date:
                    raise HoldingsError("기준 거래일을 못 정함(달력 없음)")
                rows, asof = fetch_kr(t, kr_date)
            rows = [r for r in rows if _finite(r.get("weight"))]
            if not rows:
                raise HoldingsError("보유 종목의 비중이 전부 수가 아님(NaN·무한)")
            _save(t, {"ticker": t, "market": market, "rows": rows, "asof": asof,
                      "fetched_day": today, "tried_day": today, "tries": tries, "error": "",
                      "fetched_at": now.astimezone(_KST).isoformat(timespec="seconds")})
        except Exception as exc:                               # noqa: BLE001
            # 원천이 형식을 바꿔 파서가 다른 예외(KeyError·BadZipFile·ValueError …)를 내도 **같은
            # 실패 경로**로 — 반복문을 빠져나가면 나머지 ETF 는 한 번도 안 묻고, 저장도 백오프도
            # 없이 15분마다 같은 파일을 받는다(독립 리뷰 실측: 3주기 · XLK 3콜 · 나머지 10개 0콜).
            why = (exc.reason if isinstance(exc, HoldingsError)
                   else f"보유 목록 처리 중 예외({type(exc).__name__}: {str(exc)[:80]}) — 원천 형식 "
                        "변경 또는 파서 결함")
            cur.update({"tried_day": today, "tries": tries, "error": why,
                        "tried_at": now.astimezone(_KST).isoformat(timespec="seconds")})
            try:
                _save(t, cur)
            except Exception as werr:                          # noqa: BLE001
                log.warning("etf_holdings: %s 실패 기록을 못 씀: %s", t, werr)
            res["failed"].append(t)
            log.warning("etf_holdings: %s 보유 종목 실패(오늘 %d/%d회) — %s", t, tries, MAX_TRIES, why)
            continue
        res["ok"].append(t)
    return res


def holdings_verdict(market: str, now: datetime) -> tuple[str, str]:
    """한 나라 섹터 ETF 보유 목록 — (갈래, 한 줄). 감사·`--why` 가 같이 쓴다(#35).

    fail = 목록이 하나도 없는 ETF 가 있다(그 칸은 영영 빈다 — KRX 로그인 없음 등, 우리가 고칠 것
    #303) · stale = 오늘 갱신은 실패했지만 마지막 목록이 있다 · none = 아직 한 번도 묻지 않음 ·
    ok = 전부 목록이 있다."""
    tickers = list(etfs(market))
    if not tickers:
        return "none", "섹터 ETF 목록이 없음"
    today = now.astimezone(_KST).date().isoformat()
    have, empty, old, untried = [], [], [], []
    for t in tickers:
        h = load(t) or {}
        if h.get("rows"):
            have.append(t)
            if h.get("error") and h.get("tried_day") == today and h.get("fetched_day") != today:
                old.append(t)
        elif h.get("error"):
            empty.append((t, h["error"]))
        else:
            untried.append(t)
    n = len(tickers)
    if empty:
        t, why = empty[0]
        more = f" 외 {len(empty) - 1}개" if len(empty) > 1 else ""
        return "fail", f"보유 목록이 없는 ETF {len(empty)}/{n}개 — {t}{more}: {why}"
    if len(untried) == n:
        return "none", "아직 보유 목록을 묻지 않음"
    if old:
        return "stale", f"{len(old)}/{n}개 ETF 오늘 갱신 실패 — 마지막으로 받은 목록으로 그린다"
    return "ok", f"{len(have)}/{n}개 ETF 보유 목록 있음" + (f" · 미시도 {len(untried)}" if untried else "")
